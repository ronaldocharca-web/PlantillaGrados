from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'templates' / 'acta-defensa-maestria.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)
sec.header_distance = Cm(1.25)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(10)
normal.paragraph_format.widow_control = True
title = doc.styles['Title']
title.font.name = 'Times New Roman'
title.font.size = Pt(15)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(6)
title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if style.name in ['Title', 'Normal']:
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for attr in list(fonts.attrib):
                if 'theme' in attr.lower():
                    del fonts.attrib[attr]

lang = OxmlElement('w:lang')
lang.set(qn('w:val'), 'es-BO')
normal.element.get_or_add_rPr().append(lang)
page_num = OxmlElement('w:pgNumType')
page_num.set(qn('w:start'), '36')
sec._sectPr.append(page_num)
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
hp._p.append(field)

doc.add_paragraph('ACTA DE DEFENSA', 'Title')
p = doc.add_paragraph('TESIS DE POSTGRADO MAESTRÍA')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(20)
p.runs[0].bold = True
p.runs[0].font.size = Pt(13)

def body(text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

p = body('A horas 09:10 a. m. del día viernes 10 de octubre de 2025, en modalidad presencial, con presencia del Tribunal de Defensa, se dio inicio al Acto Académico Público de Sustentación Oral de Tesis de Postgrado de Maestría titulado: ')
p.add_run('“CONDICIONES DE PARTICIPACIÓN DE LA MUJER EN EMPRENDIMIENTOS DE AGENCIAS DE VIAJE Y TURISMO EN LA CIUDAD DE LA PAZ”').bold = True
p.add_run(', elaborado por la Licenciada Karen Yamina Carvallo Toran, en el marco de la “MAESTRÍA EN DESARROLLO TURÍSTICO SUSTENTABLE”, versión ______ de la Carrera de Turismo de la Universidad Mayor de San Andrés, para optar al grado académico Magister Scientiarum (M. Sc.).')

body('Concluida la fase de exposición de la postulante, los miembros del Tribunal de Defensa formularon preguntas relacionadas con la temática de investigación, mismas que fueron respondidas por la postulante en el marco del protocolo establecido en el reglamento específico del nivel de Maestría.')

body('Finalmente, en sesión reservada y en sujeción a lo establecido en el Reglamento de Defensa de Trabajo de Investigación Postgradual de la “Universidad Boliviana”, capítulos I y II, artículos 5 al 10, el Tribunal procedió con la evaluación y calificación de la defensa oral del trabajo final, sobre la base de los siguientes criterios:')

for letter, text in [
    ('a', 'Orden de exposición.'),
    ('b', 'Ética investigativa y aportes al conocimiento de la disciplina del Turismo.'),
    ('c', 'Relación profesional-vivencial de la postulante con el tema de investigación.'),
    ('d', 'Capacidad de respuesta con rigor científico y metodológico.'),
]:
    p = doc.add_paragraph(f'{letter})  {text}')
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.space_after = Pt(4)

p = body('Habiendo decidido de manera consensuada asignar el puntaje cuantitativo y cualitativo de:')
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.keep_with_next = True
for label, value in [('NUMERAL', '92 / 100'), ('LITERAL', 'NOVENTA Y DOS'), ('HA SIDO', 'APROBADA - EXCELENTE')]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = label != 'HA SIDO'
    p.add_run(f'{label}:  ').bold = True
    p.add_run(value)

p = body('Dando legalidad a los procedimientos y resultados inscritos en el ACTA DE DEFENSA, suscriben la misma:')
p.paragraph_format.page_break_before = True
p.paragraph_format.space_after = Pt(20)

for index, (name, role) in enumerate([
    ('Dra. Juana Margot Cavero Contreras Ph. D.', 'TRIBUNAL DOCENTE'),
    ('M. Sc. Jorge Antonio Gutierrez Adauto', 'TRIBUNAL DOCENTE'),
    ('Mg. Tur. Dante Enrique Caero Miranda', 'TRIBUNAL REVISOR'),
    ('M. Sc. Lucas José Hidalgo Quezada', 'PRESIDENTE TRIBUNAL'),
]):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if index % 2:
        p.paragraph_format.left_indent = Cm(5)
    else:
        p.paragraph_format.right_indent = Cm(5)
    p.add_run('___________________________________')
    p.add_run('\n' + name)
    p.add_run('\n' + role).bold = True

doc.core_properties.title = 'Acta de defensa de tesis de postgrado de maestría'
doc.core_properties.subject = 'Maestría en Desarrollo Turístico Sustentable'
doc.core_properties.author = ''
doc.core_properties.keywords = ''
doc.save(OUT)
print(OUT)
