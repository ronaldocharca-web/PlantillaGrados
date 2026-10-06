"""Cambia el campo PAGE del encabezado por un marcador fijo de plantilla."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
archivo = ROOT / "templates" / "acta-defensa-maestria.docx"
doc = Document(archivo)
parrafo = doc.sections[0].header.paragraphs[0]

# Conserva propiedades de alineación del encabezado, pero elimina el campo PAGE
# y su valor cacheado: así las dos hojas reciben exactamente el mismo número.
for hijo in list(parrafo._p):
    if hijo.tag != qn("w:pPr"):
        parrafo._p.remove(hijo)
run = parrafo.add_run("{{pagina}}")
run.font.name = "Times New Roman"
run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
doc.save(archivo)
print(f"Número fijo configurado en las dos hojas: {archivo}")
