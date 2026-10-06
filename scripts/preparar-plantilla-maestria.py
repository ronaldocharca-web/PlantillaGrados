"""Convierte el acta transcrita a plantilla, conservando estilos, folios y firmas.

Idempotente: se puede ejecutar de nuevo sin duplicar marcadores.
"""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
archivo = ROOT / "templates" / "acta-defensa-maestria.docx"
doc = Document(archivo)
reemplazos = {
    "09:10 a. m.": "{{horaTexto}}",
    "viernes 10 de octubre de 2025": "{{fechaTexto}}",
    "CONDICIONES DE PARTICIPACIÓN DE LA MUJER EN EMPRENDIMIENTOS DE AGENCIAS DE VIAJE Y TURISMO EN LA CIUDAD DE LA PAZ": "{{tema}}",
    "la Licenciada Karen Yamina Carvallo Toran": "{{tratamiento}} {{postulante}}",
    "MAESTRÍA EN DESARROLLO TURÍSTICO SUSTENTABLE": "{{maestria}}",
    "versión ______ de": "versión   de",
    "(M. Sc.)": "({{siglasGrado}})",
    "de la postulante": "{{delPostulante}}",
    "por la postulante": "por {{elPostulante}}",
    "92 / 100": "{{notaTexto}} / 100",
    "NOVENTA Y DOS": "{{notaLiteral}}",
    "APROBADA - EXCELENTE": "{{resultado}}",
    "Dra. Juana Margot Cavero Contreras Ph. D.": "{{tribunal1}}",
    "M. Sc. Jorge Antonio Gutierrez Adauto": "{{tribunal2}}",
    "Mg. Tur. Dante Enrique Caero Miranda": "{{revisor}}",
    "M. Sc. Lucas José Hidalgo Quezada": "{{presidente}}",
}
for parrafo in doc.paragraphs:
    for run in parrafo.runs:
        texto = run.text
        for origen, destino in reemplazos.items():
            texto = texto.replace(origen, destino)
        if texto != run.text:
            run.text = texto

texto_completo = "\n".join(p.text for p in doc.paragraphs)
assert "versión   de" in texto_completo
for marcador in [v for v in reemplazos.values() if "{{" in v]:
    assert marcador in texto_completo, f"Falta el marcador {marcador}"
doc.core_properties.subject = "Plantilla variable de defensa de tesis de postgrado"
doc.save(archivo)
print(f"Plantilla actualizada: {archivo}")
